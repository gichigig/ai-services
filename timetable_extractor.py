import os
import json
import re
import io

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts text and table content from PDF bytes using PyMuPDF (fitz)."""
    text_content = []
    try:
        import fitz
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        for page_num in range(len(doc)):
            page = doc[page_num]
            # Try table extraction first if available
            try:
                tables = page.find_tables()
                if tables and tables.tables:
                    for table in tables.tables:
                        for row in table.extract():
                            text_content.append(" | ".join([str(cell or '') for cell in row]))
            except Exception:
                pass
            # Standard text extraction
            text_content.append(page.get_text())
    except Exception as e:
        print(f"Error extracting text from PDF with fitz: {e}")
    return "\n".join(text_content)

def extract_text_from_image(file_bytes: bytes, ext: str = "png") -> str:
    """Extracts text from image bytes using PyMuPDF OCR or EasyOCR."""
    text_content = []
    
    # 1. Try decoding as plain text in case text/csv file was sent
    try:
        raw = file_bytes.decode('utf-8')
        if len(raw.strip()) > 10 and not any(ord(c) < 9 or (13 < ord(c) < 32) for c in raw[:200]):
            return raw
    except Exception:
        pass

    # 2. Try PyMuPDF image support
    try:
        import fitz
        img_doc = fitz.open(stream=file_bytes, filetype=ext if ext in ["png", "jpg", "jpeg", "webp"] else "png")
        for page in img_doc:
            text = page.get_text()
            if text.strip():
                text_content.append(text)
    except Exception:
        pass

    # 3. Try EasyOCR with Spatial Table Row Reconstruction
    if not text_content or len("\n".join(text_content).strip()) < 10:
        try:
            import easyocr
            import numpy as np
            import cv2
            
            nparr = np.frombuffer(file_bytes, np.uint8)
            img = cv2.imdecode(nparr, cv2.IMREAD_COLOR)
            if img is not None:
                reader = easyocr.Reader(['en'], gpu=False)
                results = reader.readtext(img)
                
                # Spatial reconstruction of table rows:
                items = []
                heights = []
                for bbox, text, prob in results:
                    text_str = str(text).strip()
                    if not text_str:
                        continue
                    y_top = min(pt[1] for pt in bbox)
                    y_bottom = max(pt[1] for pt in bbox)
                    y_center = (y_top + y_bottom) / 2.0
                    x_left = min(pt[0] for pt in bbox)
                    h = abs(y_bottom - y_top)
                    heights.append(h)
                    items.append({
                        "y_center": y_center,
                        "x_left": x_left,
                        "height": h,
                        "text": text_str
                    })
                
                if items:
                    avg_h = (sum(heights) / len(heights)) if heights else 20.0
                    row_threshold = max(14.0, avg_h * 0.65)
                    
                    # Sort primarily by vertical position
                    items.sort(key=lambda it: it["y_center"])
                    
                    rows = []
                    curr_row = []
                    curr_y = None
                    
                    for it in items:
                        if curr_y is None or abs(it["y_center"] - curr_y) <= row_threshold:
                            curr_row.append(it)
                            if curr_y is None:
                                curr_y = it["y_center"]
                            else:
                                curr_y = (curr_y + it["y_center"]) / 2.0
                        else:
                            curr_row.sort(key=lambda it: it["x_left"])
                            rows.append(" | ".join(it["text"] for it in curr_row))
                            curr_row = [it]
                            curr_y = it["y_center"]
                            
                    if curr_row:
                        curr_row.sort(key=lambda it: it["x_left"])
                        rows.append(" | ".join(it["text"] for it in curr_row))
                        
                    text_content.append("\n".join(rows))
        except Exception as e:
            print(f"EasyOCR extraction error: {e}")

    return "\n".join(text_content)

DAYS_OF_WEEK = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

HEADER_KEYWORDS = {
    "UNIT", "UNITS", "COURSE", "COURSES", "TIME", "DAY", "DAYS", 
    "VENUE", "ROOM", "ROOMS", "LECTURER", "LECTURERS", "CODE", "CODES", 
    "TITLE", "PROGRAMME", "PROGRAM", "HOURS", "HRS"
}

def clean_ocr_time_token(token: str) -> str:
    """
    Normalizes OCR digit confusions inside time tokens (e.g.
    ZO30HRS -> 2030, IT30 -> 1730, ZOS0 -> 2030, O800 -> 0800, IIOO -> 1100,
    I4OO -> 1400, I7OO -> 1700).
    """
    t = re.sub(r'(?i)hrs?$', '', str(token or '').strip())
    t = t.strip('.:-– ')
    
    # Known standard university slot patterns
    if re.match(r'^[2Z][0ODo][S5s3][0ODo]$', t, re.I):
        return '2030'
    if re.match(r'^[1I|l][7Tt][3S5s][0ODo]$', t, re.I):
        return '1730'
    if re.match(r'^[0ODo][8B][0ODo][0ODo]$', t, re.I):
        return '0800'
    if re.match(r'^[1I|l][1I|l][0ODo][0ODo]$', t, re.I):
        return '1100'
    if re.match(r'^[1I|l][4][0ODo][0ODo]$', t, re.I):
        return '1400'
    if re.match(r'^[1I|l][7Tt][0ODo][0ODo]$', t, re.I):
        return '1700'
        
    char_map = {
        'O': '0', 'o': '0', 'D': '0',
        'I': '1', 'l': '1', 'i': '1', '|': '1', '!': '1',
        'Z': '2', 'z': '2',
        'B': '8',
    }
    return ''.join(char_map.get(c, c) for c in t)

def parse_time_slot_str(text: str):
    """Extracts start and end times in HH:MM format from headers or cells like '0800-1100', '1730-ZO30HRS', 'IT30-ZOSOHRS', '8:00 - 11:00'."""
    s = str(text or "").strip()
    
    # 1. 4-digit / OCR-confused token range (e.g. 1730-ZO30HRS, IT30-ZOSOHRS, O800-IIOOHRS)
    m_ocr = re.search(
        r'\b([0-9A-Za-z|!]{3,5})\s*(?:[-–—~]|\bto\b)\s*([0-9A-Za-z|!]{3,5}(?:HRS)?)\b',
        s,
        re.IGNORECASE
    )
    if m_ocr:
        t1 = clean_ocr_time_token(m_ocr.group(1))
        t2 = clean_ocr_time_token(m_ocr.group(2))
        if len(t1) == 4 and t1.isdigit() and len(t2) == 4 and t2.isdigit():
            return f"{t1[:2]}:{t1[2:]}", f"{t2[:2]}:{t2[2:]}"
            
    # 2. Standard colon/dot time formats (e.g. '8:00 - 11:00', '08:00-11:00', '17.30 - 20.30')
    m = re.search(r'(\d{1,2})[:.]?(\d{2})?\s*(?:[-–—~]|\bto\b)\s*(\d{1,2})[:.]?(\d{2})?', s, re.IGNORECASE)
    if m:
        s_hr, s_min, e_hr, e_min = m.group(1), m.group(2) or "00", m.group(3), m.group(4) or "00"
        return f"{int(s_hr):02d}:{s_min}", f"{int(e_hr):02d}:{e_min}"
        
    return None, None

KNOWN_COURSE_ACRONYMS = {
    "software engineering": ["BSE", "DSE"],
    "bachelor of software engineering": ["BSE"],
    "bachelor of science in software engineering": ["BSE"],
    "bsc software engineering": ["BSE"],
    "bse": ["BSE"],
    "diploma in software engineering": ["DSE"],
    "dip software engineering": ["DSE"],
    "dse": ["DSE"],
    "computer science": ["DCS", "BSCIT", "BSE"],
    "bachelor of science in information technology": ["BSCIT"],
    "information technology": ["BSCIT", "BBIT", "DIT", "DBIT", "CIT", "CICT"],
    "bachelor of business information technology": ["BBIT"],
    "diploma in information technology": ["DIT"],
    "diploma in business information technology": ["DBIT"],
    "business information technology": ["BBIT", "DBIT"],
    "bbit": ["BBIT"],
    "dbit": ["DBIT"],
    "dit": ["DIT"],
    "bscit": ["BSCIT"],
    "commerce": ["BCOM", "CBM", "DBMA"],
    "bachelor of commerce": ["BCOM"],
    "bcom": ["BCOM"],
    "business administration": ["BBAM", "DBMA", "CBM"],
    "bachelor of business administration": ["BBAM"],
    "bbam": ["BBAM"],
    "accounting and finance": ["BAF", "DAF", "DAC"],
    "accounting": ["BAF", "DAF", "DAC"],
    "baf": ["BAF"],
    "daf": ["DAF"],
    "dac": ["DAC"],
    "procurement": ["BPSM", "DPSM", "CPSM"],
    "purchasing and supplies": ["BPSM", "DPSM", "CPSM"],
    "supply chain": ["BPSM", "DPSM", "CPSM"],
    "bpsm": ["BPSM"],
    "dpsm": ["DPSM"],
    "cpsm": ["CPSM"],
    "journalism": ["BMDC", "BAJ", "DCMS", "CJMS"],
    "mass communication": ["BMDC", "BAJ", "DCMS", "CJMS"],
    "media": ["BMDC", "BAJ"],
    "hospitality": ["BHTM", "DHM", "CHTM"],
    "tourism": ["BHTM", "DTM", "CHTM"],
    "international relations": ["BIRD", "DIRD"],
    "diplomacy": ["BIRD", "DIRD"],
    "electrical engineering": ["DEE", "CEE"],
    "fashion": ["DFT"],
    "human resource": ["DHR"],
    "icdl": ["ICDL"]
}

def extract_cohort_header(text: str):
    """
    Identifies cohort and year/semester boundary headers in timetables.
    Examples:
      'BSE Y2S2' -> 'BSE Y2S2'
      'BSE Y3SI UNIT' -> 'BSE Y3S1'
      'BSE Y3S1' -> 'BSE Y3S1'
      'BBIT Y1S1' -> 'BBIT Y1S1'
      'BSE YEAR 2 SEMESTER 2' -> 'BSE Y2S2'
      'YEAR 2 SEM 2' -> 'Y2S2'
    Returns normalized cohort string if matched, otherwise None.
    """
    if not text:
        return None
    s = str(text).strip()
    if not s:
        return None

    c_upper = s.upper()
    if any(d.upper() in c_upper for d in DAYS_OF_WEEK):
        return None
    if c_upper in HEADER_KEYWORDS:
        return None

    # Normalize common OCR digit confusions in Y#S# tokens: e.g. Y3SI -> Y3S1, Y252 -> Y2S2, YIS1 -> Y1S1
    c_norm = re.sub(
        r'\b([A-Z]{2,6})\s*Y([1-4I|l])\s*[S5]\s*([1-3I|l])\b', 
        lambda m: f"{m.group(1)} Y{1 if m.group(2) in ['I', '|', 'l'] else m.group(2)}S{1 if m.group(3) in ['I', '|', 'l'] else m.group(3)}", 
        c_upper
    )

    # Pattern 1: ACRONYM + Y#S# or YEAR # SEM # (e.g. BSE Y2S2, BSE Y3S1, BBIT Y1S1, DSE Y2S2)
    m = re.search(r'\b([A-Z]{2,6})\s*(?:Y(?:EAR)?\s*([1-4])\s*S(?:EM(?:ESTER)?)?\s*([1-3])|YEAR\s*([1-4])\s*SEM(?:ESTER)?\s*([1-3]))\b', c_norm)
    if m:
        acr = m.group(1)
        y = m.group(2) or m.group(4)
        sem = m.group(3) or m.group(5)
        return f"{acr} Y{y}S{sem}"

    # Pattern 2: ACRONYM + #.# or #/# (e.g. BSE 2.2, BSE 3.1)
    m2 = re.search(r'\b([A-Z]{2,6})\s*([1-4])[\./]([1-3])\b', c_norm)
    if m2:
        return f"{m2.group(1)} Y{m2.group(2)}S{m2.group(3)}"

    # Pattern 3: Standalone Y#S# or YEAR # SEM # (e.g. Y2S2, YEAR 2 SEMESTER 2)
    m3 = re.search(r'\b(?:Y(?:EAR)?\s*([1-4])\s*S(?:EM(?:ESTER)?)?\s*([1-3])|YEAR\s*([1-4])\s*SEM(?:ESTER)?\s*([1-3]))\b', c_norm)
    if m3:
        y = m3.group(1) or m3.group(3)
        sem = m3.group(2) or m3.group(4)
        acr_match = re.search(r'\b([A-Z]{2,6})\b', c_norm)
        acr = acr_match.group(1) if acr_match and acr_match.group(1) not in ["YEAR", "SEM", "SEMESTER"] else ""
        return f"{acr + ' ' if acr else ''}Y{y}S{sem}".strip()

    # Pattern 4: Known Course Acronym + Stage / Mod / Year / Level (e.g. BSE LEVEL 2, DIT MOD 1)
    if len(c_upper.split()) <= 4:
        for phrase, acrs in KNOWN_COURSE_ACRONYMS.items():
            for acr in acrs:
                if acr in c_norm and any(k in c_norm for k in ["YEAR", "SEM", "Y1", "Y2", "Y3", "Y4", "STAGE", "MOD"]):
                    return s.strip()

    return None

def parse_day_and_time_cell(text: str):
    """
    Parses cells containing combined Day and Time slot,
    e.g. 'MONDAY 1730-ZO30HRS | EMMA WAMAITHA...' -> ('Monday', '17:30', '20:30', 'EMMA WAMAITHA...')
    or 'WEDNESDAY 1730-ZO3OHRS' -> ('Wednesday', '17:30', '20:30', '')
    """
    if not text:
        return None
    s = str(text).strip()

    day_found = None
    for d in DAYS_OF_WEEK:
        if re.search(r'\b' + re.escape(d) + r'\b', s, re.IGNORECASE) or re.search(r'\b' + re.escape(d[:3]) + r'\b', s, re.IGNORECASE):
            day_found = d
            break
            
    if not day_found:
        return None

    time_match = re.search(
        r'\b([0-9A-Za-z|!]{3,5})\s*(?:[-–—~]|\bto\b)\s*([0-9A-Za-z|!]{3,5}(?:HRS)?)\b',
        s,
        re.IGNORECASE
    )
    st, et = None, None
    time_end_pos = 0
    if time_match:
        t1 = clean_ocr_time_token(time_match.group(1))
        t2 = clean_ocr_time_token(time_match.group(2))
        if len(t1) == 4 and t1.isdigit() and len(t2) == 4 and t2.isdigit():
            st, et = f'{t1[:2]}:{t1[2:]}', f'{t2[:2]}:{t2[2:]}'
            time_end_pos = time_match.end()

    if not (st and et):
        m = re.search(r'(\d{1,2})[:.]?(\d{2})?\s*(?:[-–—~]|\bto\b)\s*(\d{1,2})[:.]?(\d{2})?', s, re.IGNORECASE)
        if m:
            s_hr, s_min, e_hr, e_min = m.group(1), m.group(2) or "00", m.group(3), m.group(4) or "00"
            st, et = f"{int(s_hr):02d}:{s_min}", f"{int(e_hr):02d}:{e_min}"
            time_end_pos = m.end()

    if st and et:
        trailing = s[time_end_pos:].strip()
        trailing = re.sub(r'^[|:\s\t-]+', '', trailing).strip()
        return day_found, st, et, trailing

    return None

def clean_unit_title(title: str) -> str:
    """Removes OCR pipe column artifacts and collapses excess whitespace while preserving full unit names and codes."""
    t = ' '.join([p.strip() for p in str(title or '').split(' | ') if p.strip()])
    t = re.sub(r'\s+', ' ', t).strip()
    return t

def parse_text_vertical_sections(raw_text: str) -> list:
    """
    Deterministically parses plain text or OCR text for vertical section timetables.
    Recognizes cohort boundary headers (e.g. 'BSE Y2S2', 'BSE Y3S1', 'BSE Y3SI UNIT') and combined Day+Time rows.
    Strictly skips blank / empty days or slots without scheduled units to prevent hallucinating classes.
    Preserves full unit names, codes, and lecturers.
    """
    blocks = []
    current_cohort = ""
    lines = [l.strip() for l in raw_text.splitlines() if l.strip()]
    i = 0
    n = len(lines)

    while i < n:
        line_str = lines[i]

        hdr = extract_cohort_header(line_str)
        if hdr:
            current_cohort = hdr
            i += 1
            continue

        dt = parse_day_and_time_cell(line_str)
        if dt:
            day, st, et, trailing = dt
            unit_parts = []

            if trailing and trailing.upper() not in HEADER_KEYWORDS and trailing not in ['-', '--', 'N/A', 'NONE', 'FREE']:
                unit_parts.append(trailing)

            # If unit_parts is empty, look ahead to the next line(s) to find the unit title
            if not unit_parts:
                while i + 1 < n:
                    next_line = lines[i + 1].strip()
                    # If next line is another Day+Time or Cohort Header, current slot is BLANK / FREE!
                    if parse_day_and_time_cell(next_line):
                        break
                    if extract_cohort_header(next_line):
                        break
                    if next_line.upper() in HEADER_KEYWORDS:
                        i += 1
                        continue
                    if next_line in ['-', '--', 'N/A', 'NONE', 'FREE']:
                        i += 1
                        break
                    unit_parts.append(next_line)
                    i += 1
                    # If next_line looks like the end of a unit description (e.g. contains venue/mode), stop consuming
                    if any(w in next_line.upper() for w in ['ONLINE', 'ROOM', 'LAB', 'HALL']):
                        break
            elif (len(trailing.split()) < 3 or trailing.rstrip().endswith(('and', 'of', '&', '-', '|'))) and i + 1 < n:
                # Incomplete trailing line, check if immediate next line continues it
                next_line = lines[i + 1].strip()
                if not parse_day_and_time_cell(next_line) and not extract_cohort_header(next_line) and next_line.upper() not in HEADER_KEYWORDS and next_line not in ['-', '--', 'N/A', 'NONE', 'FREE']:
                    unit_parts.append(next_line)
                    i += 1

            unit_title = clean_unit_title(' '.join(unit_parts))
            # Strictly check: only add if unit title is not empty, not a header keyword, and has actual content
            if unit_title and unit_title.upper() not in HEADER_KEYWORDS:
                blocks.append({
                    "title": unit_title,
                    "day_of_week": day,
                    "start_time": st,
                    "end_time": et,
                    "type": "CLASS",
                    "cohort": str(current_cohort).strip()
                })

        i += 1

    return blocks

def match_cohort_blocks(blocks: list, course: str = "", year: str = "", semester: str = "", cohort_hint: str = ""):
    """
    Intelligently matches user's course/year/semester to cohort sections in timetable.
    Strictly determines cohort boundaries and returns ONLY the blocks for the student's
    specific year and semester, preventing cross-cohort contamination.
    Returns (matched_blocks, all_cohorts, matched_cohort_name)
    """
    if not blocks:
        return [], [], None

    all_cohorts = sorted(list({b.get("cohort", "").strip() for b in blocks if b.get("cohort", "").strip()}))

    # 1. Exact hint match (e.g. user clicked "BSE Y2S2")
    if cohort_hint and cohort_hint.strip() in all_cohorts:
        target = cohort_hint.strip()
        matched = [b for b in blocks if b.get("cohort", "").strip() == target]
        return matched, all_cohorts, target

    # If only one cohort exists in the document, return all blocks
    if len(all_cohorts) <= 1:
        cohort_name = all_cohorts[0] if all_cohorts else None
        return blocks, all_cohorts, cohort_name

    combined_text = f"{course} {year} {semester} {cohort_hint}".strip().lower()
    if not combined_text:
        # Multiple cohorts exist, but no user profile or hint was given:
        # Return only the first cohort's blocks to avoid combining multiple years!
        first_cohort = all_cohorts[0]
        matched = [b for b in blocks if b.get("cohort", "").strip() == first_cohort]
        return matched, all_cohorts, None

    y_match = re.search(r"(?:year|y|yr|level|mod|module)\s*(\d)", combined_text) or re.search(r"\b([1-4])\b", year or "")
    y_num = y_match.group(1) if y_match else ""

    s_match = re.search(r"(?:semester|sem|s|term|t)\s*(\d)", combined_text) or re.search(r"\b([1-4])\b", semester or "")
    s_num = s_match.group(1) if s_match else ""

    level = ""
    if any(k in combined_text for k in ["bachelor", "degree", "undergraduate", "bsc", "b.sc", "ba "]):
        level = "B"
    elif any(k in combined_text for k in ["diploma", "dip", "dip."]):
        level = "D"
    elif any(k in combined_text for k in ["certificate", "cert", "cert."]):
        level = "C"

    candidates = []
    for phrase, acronyms in KNOWN_COURSE_ACRONYMS.items():
        if phrase in combined_text or phrase.replace(" ", "") in combined_text.replace(" ", ""):
            for acr in acronyms:
                if level and acr[0] != level:
                    continue
                candidates.append(acr)

    if not candidates:
        words = re.findall(r"[a-zA-Z]+", str(course).upper())
        for w in words:
            if len(w) in [3, 4, 5]:
                candidates.append(w)

    candidates = list(dict.fromkeys(candidates))

    best_cohort = None
    best_score = -999

    for c in all_cohorts:
        c_upper = c.upper()
        score = 0

        # Course acronym match
        matched_cand = False
        for cand in candidates:
            if cand in c_upper:
                score += 20
                if c_upper.startswith(cand):
                    score += 10
                matched_cand = True
                break
        if candidates and not matched_cand:
            score -= 15

        # Year match / penalty
        c_y_match = re.search(r"(?:Y|YEAR|MOD)\s*([1-4])", c_upper)
        c_y_num = c_y_match.group(1) if c_y_match else ""
        if y_num:
            if c_y_num:
                if c_y_num == y_num:
                    score += 25
                else:
                    # STRICT PENALTY: Year mismatch! (e.g. Student is Y2, cohort is Y3)
                    score -= 50
            elif f"Y{y_num}" in c_upper or f"YEAR {y_num}" in c_upper:
                score += 25
            else:
                score -= 20

        # Semester match / penalty
        c_s_match = re.search(r"(?:S|SEM|SEMESTER|TERM)\s*([1-4])", c_upper)
        c_s_num = c_s_match.group(1) if c_s_match else ""
        if s_num:
            if c_s_num:
                if c_s_num == s_num:
                    score += 25
                else:
                    # STRICT PENALTY: Semester mismatch! (e.g. Student is S2, cohort is S1)
                    score -= 30
            elif f"S{s_num}" in c_upper or f"SEM {s_num}" in c_upper:
                score += 25
            else:
                score -= 15

        if score > best_score:
            best_score = score
            best_cohort = c

    print(f"Cohort Matching: best_cohort='{best_cohort}', score={best_score}, all={all_cohorts}")

    if best_cohort and best_score >= 15:
        # STRICT FILTERING: Only return blocks belonging EXACTLY to the best cohort!
        matched = [b for b in blocks if b.get("cohort", "").strip() == best_cohort]
        return matched, all_cohorts, best_cohort

    first_cohort = all_cohorts[0] if all_cohorts else None
    matched = [b for b in blocks if b.get("cohort", "").strip() == first_cohort] if first_cohort else blocks
    return matched, all_cohorts, None

def parse_excel_grid(file_bytes: bytes, ext: str = "xlsx") -> list:
    """
    Directly and deterministically parses Excel timetable spreadsheets:
    - Vertical Section Timetables: Column A has combined Day+Time (e.g. 'MONDAY 1730-2030HRS'),
      Column B has lecturer/unit info, separated by cohort header rows ('BSE Y2S2', 'BSE Y3S1').
    - Horizontal Grid Timetables: Column headers have times (e.g. 0800-1100), rows have days.
    - Accurately establishes section boundaries between cohorts.
    """
    blocks = []
    
    # 1. Handle CSV format
    if ext == "csv":
        try:
            import csv
            text = file_bytes.decode('utf-8', errors='replace')
            reader = csv.reader(io.StringIO(text))
            current_cohort = "Timetable"
            for row in reader:
                non_empty = [c for c in row if c is not None and str(c).strip()]
                if not non_empty:
                    continue
                first_str = str(non_empty[0]).strip()
                detected = extract_cohort_header(first_str)
                if not detected and len(non_empty) <= 3:
                    detected = extract_cohort_header(" ".join(str(c).strip() for c in non_empty))
                if detected:
                    current_cohort = detected
                    continue
                dt = parse_day_and_time_cell(first_str)
                if dt:
                    day, st, et = dt[0], dt[1], dt[2]
                    trailing = dt[3] if len(dt) > 3 else ""
                    unit_parts = [str(c).strip() for c in non_empty[1:] if c and str(c).strip()]
                    if not unit_parts and trailing:
                        unit_parts = [trailing]
                    if unit_parts:
                        unit_text = " ".join(unit_parts).lstrip("| \t")
                        if unit_text.upper() not in ["UNIT", "UNITS", "-", "--", "N/A", "NONE", "FREE"]:
                            blocks.append({
                                "title": unit_text,
                                "day_of_week": day,
                                "start_time": st,
                                "end_time": et,
                                "type": "CLASS",
                                "cohort": str(current_cohort).strip()
                            })
            return blocks
        except Exception as csv_err:
            print(f"Error parsing CSV in parse_excel_grid: {csv_err}")

    # 2. Handle XLSX / XLS format
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            current_cohort = sheet_name
            current_col_times = {}

            for row in sheet.iter_rows(values_only=True):
                non_empty = [c for c in row if c is not None and str(c).strip()]
                if not non_empty:
                    continue

                first_non_empty_str = str(non_empty[0]).strip()

                # A. Check if row is a Cohort / Section Header (e.g. 'BSE Y2S2', 'BSE Y3S1')
                detected_cohort = extract_cohort_header(first_non_empty_str)
                if not detected_cohort and len(non_empty) <= 3:
                    detected_cohort = extract_cohort_header(" ".join(str(c).strip() for c in non_empty))

                if detected_cohort:
                    current_cohort = detected_cohort
                    current_col_times = {} # Reset grid state when a new cohort section starts
                    continue

                # B. Check if row is a Vertical Day+Time row (Col A has combined Day+Time, Col B has unit)
                # E.g. 'MONDAY 1730-2030HRS' | 'EMMA WAMAITHA BSD 222 Design and analysis and algorithm ONLINE'
                dt_match = parse_day_and_time_cell(first_non_empty_str)
                if dt_match:
                    day, st, et = dt_match[0], dt_match[1], dt_match[2]
                    trailing = dt_match[3] if len(dt_match) > 3 else ""
                    unit_cells = [str(c).strip() for c in non_empty[1:] if c and str(c).strip()]
                    if not unit_cells and trailing:
                        unit_cells = [trailing]
                    if unit_cells:
                        unit_text = " ".join(unit_cells).lstrip("| \t")
                        if unit_text.upper() not in ["UNIT", "UNITS", "-", "--", "N/A", "NONE", "FREE"]:
                            blocks.append({
                                "title": unit_text,
                                "day_of_week": day,
                                "start_time": st,
                                "end_time": et,
                                "type": "CLASS",
                                "cohort": str(current_cohort).strip()
                            })
                    continue

                # C. Check if Col 0 is Day and Col 1 is Time
                if len(non_empty) >= 2:
                    first_day = next((d for d in DAYS_OF_WEEK if d.lower() == first_non_empty_str.lower() or first_non_empty_str.lower().startswith(d.lower()[:3])), None)
                    if first_day:
                        second_str = str(non_empty[1]).strip()
                        st, et = parse_time_slot_str(second_str)
                        if st and et:
                            unit_parts = [str(c).strip() for c in non_empty[2:] if c and str(c).strip()]
                            if unit_parts:
                                unit_text = " ".join(unit_parts).lstrip("| \t")
                                if unit_text.upper() not in ["UNIT", "UNITS", "-", "--", "N/A", "NONE", "FREE"]:
                                    blocks.append({
                                        "title": unit_text,
                                        "day_of_week": first_day,
                                        "start_time": st,
                                        "end_time": et,
                                        "type": "CLASS",
                                        "cohort": str(current_cohort).strip()
                                    })
                            continue

                # D. Check if row is a Horizontal Grid Day row
                day_match = None
                for cell in row[:3]:
                    if cell:
                        cell_clean = str(cell).strip().lower()
                        for d in DAYS_OF_WEEK:
                            if d.lower() == cell_clean or cell_clean.startswith(d.lower()[:3]):
                                day_match = d
                                break
                    if day_match:
                        break

                if day_match and current_col_times:
                    for col_idx, (st, et) in current_col_times.items():
                        if col_idx < len(row):
                            cell_val = row[col_idx]
                            if cell_val and str(cell_val).strip():
                                val_str = str(cell_val).strip()
                                if val_str.lower() not in [d.lower() for d in DAYS_OF_WEEK] and not parse_time_slot_str(val_str)[0]:
                                    if val_str.upper() not in ["UNIT", "UNITS", "-", "--", "N/A", "NONE", "FREE"]:
                                        blocks.append({
                                            "title": val_str,
                                            "day_of_week": day_match,
                                            "start_time": st,
                                            "end_time": et,
                                            "type": "CLASS",
                                            "cohort": str(current_cohort).strip()
                                        })
                    continue

                # E. Check if row is a Horizontal Grid Time Header row
                time_cols = {}
                for col_idx, cell in enumerate(row):
                    if cell:
                        st, et = parse_time_slot_str(cell)
                        if st and et:
                            time_cols[col_idx] = (st, et)
                if len(time_cols) >= 2:
                    current_col_times = time_cols
                    continue

                # F. Fallback section header for rows with few cells
                if len(non_empty) <= 3:
                    hdr_text = " ".join(str(c).strip() for c in non_empty)
                    if not any(d.lower() in hdr_text.lower() for d in DAYS_OF_WEEK) and not parse_time_slot_str(hdr_text)[0]:
                        if hdr_text.upper() not in ["UNIT", "UNITS", "DAY", "TIME"]:
                            current_cohort = hdr_text

    except Exception as e:
        print(f"Error in parse_excel_grid: {e}")

    return blocks

def clean_and_parse_json(raw_response: str) -> list:
    """Cleans markdown wrappers or explanations and parses JSON list of blocks."""
    text = raw_response.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.MULTILINE)
        text = re.sub(r"```\s*$", "", text, flags=re.MULTILINE)
        text = text.strip()

    match = re.search(r"\[\s*\{.*\}\s*\]", text, re.DOTALL)
    if match:
        text = match.group(0)

    try:
        parsed = json.loads(text)
        if isinstance(parsed, list):
            valid_blocks = []
            for item in parsed:
                if isinstance(item, dict) and "title" in item and "day_of_week" in item:
                    title = str(item.get("title", "")).strip()
                    # Strictly discard empty or hallucinated placeholder titles
                    if not title or title.upper() in ["LECTURE SESSION", "LECTURE", "CLASS", "STUDY", "FREE", "N/A", "NONE", "-", "--", "TBD"]:
                        continue

                    day = str(item.get("day_of_week", "Monday")).capitalize()
                    valid_days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
                    matched_day = next((d for d in valid_days if d.lower() in day.lower()), None)
                    if not matched_day:
                        continue
                    
                    st = str(item.get("start_time", "08:00")).strip()
                    et = str(item.get("end_time", "10:00")).strip()

                    b_type = str(item.get("type", "CLASS")).upper()
                    if b_type not in ["CLASS", "STUDY"]:
                        b_type = "CLASS"
                    
                    valid_blocks.append({
                        "title": title,
                        "day_of_week": matched_day,
                        "start_time": st,
                        "end_time": et,
                        "type": b_type,
                        "cohort": str(item.get("cohort", "")).strip()
                    })
            return valid_blocks
    except Exception as e:
        print(f"Failed to parse JSON timetable blocks: {e}\nRaw output was: {raw_response[:300]}")
    
    return []

async def extract_timetable_from_bytes(
    file_bytes: bytes, 
    filename: str = "document.pdf", 
    course: str = "",
    year: str = "",
    semester: str = "",
    cohort_hint: str = ""
) -> dict:
    """
    Main extraction pipeline:
    1. For Excel files (.xlsx, .xls, .csv): Deterministically parses grid tables and vertical section timetables.
    2. For text/pdf/image: Deterministically parses section boundaries first, falling back to LLM.
    3. Runs smart cohort matcher to strictly isolate student's exact year and semester.
    4. Returns dict containing matched blocks, all_cohorts list, and matched_cohort name.
    """
    ext = filename.split(".")[-1].lower() if "." in filename else "pdf"
    
    # 1. Native Excel Parser (Vertical sections & 2D Grid)
    if ext in ["xlsx", "xls", "csv"]:
        all_blocks = parse_excel_grid(file_bytes, ext=ext)
        if all_blocks:
            matched_blocks, all_cohorts, matched_cohort = match_cohort_blocks(
                all_blocks, course=course, year=year, semester=semester, cohort_hint=cohort_hint
            )
            print(f"Excel Extraction: {len(all_blocks)} total blocks, matched '{matched_cohort}' ({len(matched_blocks)} blocks), {len(all_cohorts)} cohorts found.")
            return {
                "blocks": matched_blocks,
                "all_cohorts": all_cohorts,
                "matched_cohort": matched_cohort,
                "all_blocks": all_blocks
            }

    # 2. Extract raw text from PDF / Image / Text
    if ext in ["pdf"]:
        raw_text = extract_text_from_pdf(file_bytes)
    elif ext in ["txt"]:
        try:
            raw_text = file_bytes.decode('utf-8', errors='ignore')
        except Exception:
            raw_text = extract_text_from_image(file_bytes, ext)
    else:
        raw_text = extract_text_from_image(file_bytes, ext)

    if not raw_text.strip():
        print("No text could be extracted from the uploaded document.")
        return {"blocks": [], "all_cohorts": [], "matched_cohort": None, "all_blocks": []}

    # 3. Attempt deterministic text parsing for vertical section structures
    text_blocks = parse_text_vertical_sections(raw_text)
    if text_blocks and len(text_blocks) >= 1:
        matched_blocks, all_cohorts, matched_cohort = match_cohort_blocks(
            text_blocks, course=course, year=year, semester=semester, cohort_hint=cohort_hint
        )
        print(f"Deterministic Text Extraction: {len(text_blocks)} blocks, matched '{matched_cohort}' ({len(matched_blocks)} blocks), {len(all_cohorts)} cohorts found.")
        return {
            "blocks": matched_blocks,
            "all_cohorts": all_cohorts,
            "matched_cohort": matched_cohort,
            "all_blocks": text_blocks
        }

    # 4. LLM Extraction Fallback
    target_info = ""
    if course or year or semester or cohort_hint:
        target_info = f"Target Student Info: Course='{course}', Year='{year}', Semester='{semester}', Cohort='{cohort_hint}'.\n"

    system_prompt = (
        "You are an expert academic timetable extraction engine. "
        "Your task is to parse raw document/timetable text, identify distinct cohort/year/semester sections, "
        "and output ONLY a valid JSON array of lecture or study blocks. "
        "Do not include conversational greetings, explanations, markdown formatting, or any text other than the JSON array."
    )

    user_prompt = f"""
{target_info}Extract distinct class, lesson, lecture, laboratory, tutorial, and study blocks from the following timetable document.

DOCUMENT CONTENT:
\"\"\"
{raw_text[:12000]}
\"\"\"

CRITICAL EXTRACTION RULES:
1. COHORT BOUNDARIES: The document may contain multiple classes/cohorts separated by section headers (e.g. 'BSE Y2S2', 'BSE Y3S1', 'Year 2 Semester 2', 'BCOM Y1S1').
   - Each section begins with its header (e.g. 'BSE Y2S2') and ENDS when the next section header (e.g. 'BSE Y3S1') appears.
   - For every block, you MUST populate the 'cohort' field with the exact cohort/section header it belongs to (e.g. 'BSE Y2S2').
   - NEVER combine or mix blocks across different years or semesters.
2. If Target Student Info is specified above (e.g. Year 2 Semester 2), extract ONLY the classes belonging to that specific year and semester, stopping at the boundary where that year and semester section ends. Do NOT include classes from other years or semesters.
3. BLANK / EMPTY DAYS AND TIME SLOTS (CRITICAL - NO HALLUCINATION):
   - In timetables, many days or time slots are empty/blank with NO class scheduled (for example, Wednesday or Saturday early mornings or late evenings with no subject listed).
   - If a day or time slot has no unit title or is followed by another day/time slot, it is a BLANK slot.
   - NEVER invent, hallucinate, or assume units or subjects for blank days or empty time slots!
   - DO NOT shift a class from Thursday into Wednesday! If Wednesday is blank, DO NOT output any block for Wednesday.
   - Only output blocks where an actual unit/course code or subject title is explicitly printed.
4. 'day_of_week' MUST be one of: "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday".
5. 'start_time' and 'end_time' MUST be in 24-hour "HH:MM" format (e.g. "08:00", "09:30", "17:30", "20:30"). If hours are written like '1730-2030HRS', format them as '17:30' and '20:30'.
6. 'title' should include lecturer name, unit/course code, unit title, and mode/venue (e.g. "EMMA WAMAITHA BSD 222 Design and analysis and algorithm ONLINE").
7. Return ONLY a valid JSON array of objects with keys: "title", "day_of_week", "start_time", "end_time", "cohort".

Format example:
[
  {{"title": "EMMA WAMAITHA BSD 222 Design and analysis and algorithm ONLINE", "day_of_week": "Monday", "start_time": "17:30", "end_time": "20:30", "cohort": "BSE Y2S2"}},
  {{"title": "GABRIEL MUCHANGI BSD 226 Software Metrics ONLINE", "day_of_week": "Tuesday", "start_time": "17:30", "end_time": "20:30", "cohort": "BSE Y2S2"}}
]
"""

    import modal
    blocks = []
    
    # Try Small Model directly
    try:
        SmallModelCls = modal.Cls.from_name("bruv-schools-llm", "QwenSmallModel")
        small_model = SmallModelCls()
        response = await small_model.generate.remote.aio(system_prompt, user_prompt)
        blocks = clean_and_parse_json(response)
    except Exception as e:
        print(f"Small model extraction error: {e}")

    # Fallback to Large Model if needed
    if not blocks:
        try:
            print("Retrying timetable extraction with QwenLargeModel...")
            LargeModelCls = modal.Cls.from_name("bruv-schools-llm", "QwenLargeModel")
            large_model = LargeModelCls()
            response = await large_model.generate.remote.aio(system_prompt, user_prompt)
            blocks = clean_and_parse_json(response)
        except Exception as e:
            print(f"Large model extraction error: {e}")

    matched_blocks, all_cohorts, matched_cohort = match_cohort_blocks(
        blocks, course=course, year=year, semester=semester, cohort_hint=cohort_hint
    )
    return {
        "blocks": matched_blocks,
        "all_cohorts": all_cohorts,
        "matched_cohort": matched_cohort,
        "all_blocks": blocks
    }

async def generate_study_timetable_ai(class_blocks: list, course: str = "", study_mode: str = "FULL_TIME") -> list:
    """
    Uses Qwen to generate an optimized weekly study & revision timetable
    fitting around the student's scheduled class lecture blocks.
    """
    import modal
    
    classes_summary = "\n".join([
        f"- {b.get('day_of_week')}: {b.get('start_time')}-{b.get('end_time')} : {b.get('title')}"
        for b in class_blocks
    ])
    
    system_prompt = (
        "You are an expert academic tutor and study planner for university students. "
        "Your task is to generate a realistic, balanced, high-yield weekly self-study timetable. "
        "Output ONLY a valid JSON array of study session blocks. "
        "Do not include any other text, greetings, markdown codeblocks or explanations."
    )
    
    mode_text = "Part-time student (prefer evenings 18:00-21:00 and weekends)" if "PART" in str(study_mode).upper() else (
        "Online/Distance student (flexible schedule throughout the week)" if "ONLINE" in str(study_mode).upper() else
        "Full-time student (prefer weekday late afternoons/evenings 17:30-20:30 and weekend mornings)"
    )
    
    user_prompt = f"""
Student Course: {course or 'University Degree'}
Study Mode: {mode_text}

STUDENT'S EXISTING CLASS / LECTURE SCHEDULE (DO NOT SCHEDULE STUDY SESSIONS DURING THESE TIMES):
\"\"\"
{classes_summary if classes_summary else 'No fixed classes provided (create full balanced study schedule across core subjects)'}
\"\"\"

Instructions:
1. Create 6 to 10 study/revision sessions across the week (Monday through Sunday).
2. Schedule sessions during free hours (e.g. 17:30-19:30, 19:00-21:00, or weekend mornings 09:00-12:00) without overlapping any lecture above.
3. Each block 'title' must specify the subject focus and goal (e.g. "Review Structured Programming: Loops & Functions", "Math 111: Calculus Practice Problems", "Digital Literacy: Lab Assignment Prep").
4. 'day_of_week' must be: "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", or "Sunday".
5. 'start_time' and 'end_time' in 24-hour "HH:MM" format.
6. 'type' MUST be "STUDY".
7. Return ONLY a valid JSON array of objects with keys: "title", "day_of_week", "start_time", "end_time", "type".
"""

    study_blocks = []
    try:
        SmallModelCls = modal.Cls.from_name("bruv-schools-llm", "QwenSmallModel")
        small_model = SmallModelCls()
        resp = await small_model.generate.remote.aio(system_prompt, user_prompt)
        study_blocks = clean_and_parse_json(resp)
    except Exception as e:
        print(f"Small model study generation error: {e}")
        
    if not study_blocks:
        try:
            LargeModelCls = modal.Cls.from_name("bruv-schools-llm", "QwenLargeModel")
            large_model = LargeModelCls()
            resp = await large_model.generate.remote.aio(system_prompt, user_prompt)
            study_blocks = clean_and_parse_json(resp)
        except Exception as e:
            print(f"Large model study generation error: {e}")
            
    for b in study_blocks:
        b["type"] = "STUDY"
        
    return study_blocks

